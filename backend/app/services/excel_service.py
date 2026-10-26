import os
import hashlib
from pathlib import Path
from datetime import datetime, timezone
from typing import Optional, Dict, Any, List, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.session import ExamSession
from app.models.student import Student
from app.models.subject import Subject
from app.models.result import StudentResult, SubjectResult
from app.models.excel_report import ExcelReport
from app.models.job import AutomationJobLog
from app.services.result_service import get_result_matrix, PASSING_GRADES, GRADE_POINTS

# Enterprise Institutional Color Palette
NAVY_HEADER = "0F172A"       # Deep Slate 900
NAVY_SUBHEADER = "1E293B"    # Slate 800
ACCENT_BLUE = "2563EB"       # Royal Blue
ZEBRA_ROW = "F8FAFC"         # Clean Light Grey / Slate 50
CARD_BG = "F1F5F9"           # Slate 100
BORDER_COLOR = "000000"      # Standard crisp black border (True All Borders)
BORDER_DARK = "000000"

PASS_FILL = "DCFCE7"         # Green 100
PASS_FONT = "15803D"         # Green 700
FAIL_FILL = "FEE2E2"         # Red 100
FAIL_FONT = "B91C1C"         # Red 700
AMBER_FILL = "FEF3C7"        # Amber 100
AMBER_FONT = "B45309"        # Amber 700
INFO_FILL = "E0F2FE"         # Sky 100
INFO_FONT = "0369A1"         # Sky 700
MUTED_FILL = "F1F5F9"        # Slate 100
MUTED_FONT = "64748B"        # Slate 500

font_title = Font(name="Times New Roman", size=14, bold=True, color="FFFFFF")
font_subtitle = Font(name="Times New Roman", size=10, italic=True, color="CBD5E1")
font_header = Font(name="Times New Roman", size=10, bold=True, color="FFFFFF")
font_sub_header = Font(name="Times New Roman", size=9, bold=False, color="E2E8F0")
font_data = Font(name="Times New Roman", size=10, color="000000")
font_bold_data = Font(name="Times New Roman", size=10, bold=True, color="000000")
font_kpi_num = Font(name="Times New Roman", size=14, bold=True, color="0F172A")
font_kpi_label = Font(name="Times New Roman", size=9, bold=True, color="64748B")
font_summary_total = Font(name="Times New Roman", size=10, bold=True, color="0F172A")

fill_header = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
fill_sub_header = PatternFill(start_color=NAVY_SUBHEADER, end_color=NAVY_SUBHEADER, fill_type="solid")
fill_zebra = PatternFill(start_color=ZEBRA_ROW, end_color=ZEBRA_ROW, fill_type="solid")
fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
fill_kpi = PatternFill(start_color=CARD_BG, end_color=CARD_BG, fill_type="solid")
fill_summary_total = PatternFill(start_color=CARD_BG, end_color=CARD_BG, fill_type="solid")

thin_side = Side(border_style="thin", color=BORDER_COLOR)
table_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
summary_border = Border(
    left=thin_side,
    right=thin_side,
    top=Side(border_style="thin", color=BORDER_DARK),
    bottom=Side(border_style="double", color="0F172A")
)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)


def apply_sheet_view_and_freeze(ws, freeze_cell: Optional[str] = "C6") -> None:
    """
    Configures grid lines, ensures no unintended worksheet protection or scroll restrictions,
    and sets proper freeze panes with clean active pane selections to prevent jumping/locking.
    """
    if ws.views.sheetView:
        ws.views.sheetView[0].showGridLines = True
    else:
        ws.sheet_view.showGridLines = True

    # Remove any worksheet protection or scroll restrictions
    ws.protection.sheet = False
    ws.protection.disable()

    if freeze_cell:
        ws.freeze_panes = freeze_cell
        # Fix openpyxl selection on active pane to match the top-left cell of the scrollable region
        if ws.sheet_view and ws.sheet_view.selection:
            for s in ws.sheet_view.selection:
                if s.pane == "bottomRight":
                    s.activeCell = freeze_cell
                    s.sqref = freeze_cell


def apply_print_settings(ws, landscape: bool = True, print_title_row: int = 5) -> None:
    """
    Applies standard enterprise print configuration: repeated headers, A4 sizing, fit to page width.
    """
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE if landscape else ws.ORIENTATION_PORTRAIT
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = f"${print_title_row}:${print_title_row}"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0

    ws.page_margins.left = 0.5
    ws.page_margins.right = 0.5
    ws.page_margins.top = 0.75
    ws.page_margins.bottom = 0.75

    ws.oddFooter.center.text = "Page &P of &N"
    ws.oddFooter.right.text = "&D &T"
    ws.oddHeader.left.text = settings.INSTITUTION_NAME


def apply_auto_column_widths(
    ws,
    header_row: int = 5,
    max_data_row: Optional[int] = None,
    min_col: int = 1,
    max_col: Optional[int] = None,
    compact_cols: Optional[List[int]] = None,
    max_width_cap: int = 55
) -> None:
    """
    Calculates dynamic column widths starting from the table header row.
    Ignores Excel formula strings and multi-column merged summary titles to prevent
    artificially inflated column widths. Compacts subject and grade columns.
    """
    if not max_col:
        max_col = ws.max_column
    if not max_data_row:
        max_data_row = ws.max_row

    compact_set = set(compact_cols or [])

    for col_idx in range(min_col, max_col + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0

        # Scan from header_row down to max_data_row
        for row in range(header_row, max_data_row + 1):
            cell = ws.cell(row=row, column=col_idx)
            val = cell.value
            if val is not None:
                val_str = str(val).strip()
                # Skip formulas starting with '=' so formula strings don't expand column width
                if val_str.startswith("="):
                    continue
                first_line = val_str.split("\n")[0]
                if len(first_line) > max_len:
                    max_len = len(first_line)

        if col_idx in compact_set:
            ws.column_dimensions[col_letter].width = max(max_len + 3, 10)
        elif col_idx == 1:
            ws.column_dimensions[col_letter].width = max(max_len + 3, 7)
        else:
            calculated = max_len + 4
            ws.column_dimensions[col_letter].width = min(max(calculated, 11), max_width_cap)


def apply_conditional_formatting_rules(ws, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
    """
    Applies live Excel Conditional Formatting rules across the dynamic grade & status cells.
    """
    if end_row < start_row or end_col < start_col:
        return

    cell_range = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"

    # Pass rule: green fill and text
    pass_rule = CellIsRule(
        operator="equal",
        formula=['"PASS"'],
        stopIfTrue=False,
        fill=PatternFill(start_color=PASS_FILL, end_color=PASS_FILL, fill_type="solid"),
        font=Font(name="Times New Roman", size=10, bold=True, color=PASS_FONT)
    )
    ws.conditional_formatting.add(cell_range, pass_rule)

    # Fail rule: red fill and text
    fail_rule = CellIsRule(
        operator="equal",
        formula=['"FAIL"'],
        stopIfTrue=False,
        fill=PatternFill(start_color=FAIL_FILL, end_color=FAIL_FILL, fill_type="solid"),
        font=Font(name="Times New Roman", size=10, bold=True, color=FAIL_FONT)
    )
    ws.conditional_formatting.add(cell_range, fail_rule)

    # Arrear grades: RA, SA, AB, W
    for ar_grade in ["RA", "SA", "AB", "W"]:
        ar_rule = CellIsRule(
            operator="equal",
            formula=[f'"{ar_grade}"'],
            stopIfTrue=False,
            fill=PatternFill(start_color=FAIL_FILL, end_color=FAIL_FILL, fill_type="solid"),
            font=Font(name="Times New Roman", size=10, bold=True, color=FAIL_FONT)
        )
        ws.conditional_formatting.add(cell_range, ar_rule)

    # Passing grades: O, A+, A, B+, B, C
    for p_grade in ["O", "A+", "A", "B+", "B", "C"]:
        p_rule = CellIsRule(
            operator="equal",
            formula=[f'"{p_grade}"'],
            stopIfTrue=False,
            fill=PatternFill(start_color=PASS_FILL, end_color=PASS_FILL, fill_type="solid"),
            font=Font(name="Times New Roman", size=10, bold=True, color=PASS_FONT)
        )
        ws.conditional_formatting.add(cell_range, p_rule)


def build_result_analysis_sheet(ws, matrix):
    ws.title = "Result Analysis"

    # Base headers + dynamic subjects + academic metrics
    headers = ["S.No", "Register No", "Student Name", "Branch", "Batch", "Sem"]
    sub_col_start = len(headers) + 1

    for col in matrix.columns:
        headers.append(col.code)

    headers.extend(["Current Arrears", "Standing Arrears", "GPA", "Result Status", "Verified"])
    total_cols = len(headers)
    last_col_letter = get_column_letter(total_cols)

    # Title Block Banner (Rows 1 to 3) — dynamically merged across all columns
    ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = settings.INSTITUTION_NAME.upper()
    ws["A1"].font = font_title
    ws["A1"].fill = fill_header
    ws["A1"].alignment = align_center

    ws.merge_cells(f"A2:{last_col_letter}2")
    ws["A2"] = f"OFFICE OF CONTROLLER OF EXAMINATIONS — COMPREHENSIVE RESULT ANALYSIS ({matrix.session_code})"
    ws["A2"].font = font_header
    ws["A2"].fill = fill_header
    ws["A2"].alignment = align_center

    ws.merge_cells(f"A3:{last_col_letter}3")
    meta_info = (
        f"Department: {matrix.department_code or 'ALL'} | "
        f"Total Candidates: {matrix.total_students} | "
        f"Pass Rate: {matrix.pass_percentage}% | "
        f"Avg GPA: {matrix.average_gpa} | "
        f"Exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"
    )
    ws["A3"] = meta_info
    ws["A3"].font = font_subtitle
    ws["A3"].fill = fill_header
    ws["A3"].alignment = align_center

    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 10  # Visual buffer row

    # Row 5: Table Header
    header_row = 5
    ws.row_dimensions[header_row].height = 28

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = table_border

    # Data Rows (Row 6 onwards)
    data_start_row = 6
    compact_cols = []
    num_subjects = len(matrix.columns)
    for s_idx in range(num_subjects):
        compact_cols.append(sub_col_start + s_idx)

    curr_arr_col = sub_col_start + num_subjects
    stand_arr_col = curr_arr_col + 1
    gpa_col = stand_arr_col + 1
    status_col = gpa_col + 1
    verif_col = status_col + 1

    last_data_row = data_start_row - 1

    for i, student in enumerate(matrix.rows, 1):
        r = data_start_row + i - 1
        last_data_row = r
        ws.row_dimensions[r].height = 20
        row_fill = fill_zebra if i % 2 == 0 else fill_white

        # Base Columns
        cell_sno = ws.cell(row=r, column=1, value=i)
        cell_sno.alignment = align_center
        cell_sno.font = font_data

        reg_cell = ws.cell(row=r, column=2, value=str(student.register_number).strip())
        reg_cell.number_format = "@"
        reg_cell.alignment = align_center
        reg_cell.font = font_data

        name_cell = ws.cell(row=r, column=3, value=student.student_name)
        name_cell.alignment = align_left
        name_cell.font = font_data

        dept_cell = ws.cell(row=r, column=4, value=student.department_code)
        dept_cell.alignment = align_center
        dept_cell.font = font_data

        batch_cell = ws.cell(row=r, column=5, value=student.batch)
        batch_cell.alignment = align_center
        batch_cell.font = font_data

        sem_cell = ws.cell(row=r, column=6, value=student.semester)
        sem_cell.alignment = align_center
        sem_cell.font = font_data

        # Subject Grades
        for s_idx, sub in enumerate(matrix.columns):
            c_idx = sub_col_start + s_idx
            grade = student.grades.get(sub.code, "-")
            grade_cell = ws.cell(row=r, column=c_idx, value=grade)
            grade_cell.alignment = align_center

            if grade in ["RA", "SA", "AB", "W"]:
                grade_cell.font = Font(name="Times New Roman", size=10, bold=True, color=FAIL_FONT)
                grade_cell.fill = PatternFill(start_color=FAIL_FILL, end_color=FAIL_FILL, fill_type="solid")
            elif grade in PASSING_GRADES:
                grade_cell.font = Font(name="Times New Roman", size=10, bold=True, color=PASS_FONT)
                grade_cell.fill = PatternFill(start_color=PASS_FILL, end_color=PASS_FILL, fill_type="solid")
            else:
                grade_cell.font = Font(name="Times New Roman", size=10, color=MUTED_FONT)
                grade_cell.fill = PatternFill(start_color=MUTED_FILL, end_color=MUTED_FILL, fill_type="solid")

        # Academic Totals & Status
        curr_arr_cell = ws.cell(row=r, column=curr_arr_col, value=student.arrears_count)
        curr_arr_cell.alignment = align_center
        curr_arr_cell.font = font_data

        stand_arr_cell = ws.cell(row=r, column=stand_arr_col, value=student.standing_arrears_count)
        stand_arr_cell.alignment = align_center
        stand_arr_cell.font = font_data

        gpa_cell = ws.cell(row=r, column=gpa_col, value=student.gpa)
        gpa_cell.number_format = "0.00"
        gpa_cell.alignment = align_center
        gpa_cell.font = font_data

        status_cell = ws.cell(row=r, column=status_col, value=student.result_status)
        status_cell.alignment = align_center
        if student.result_status == "PASS":
            status_cell.fill = PatternFill(start_color=PASS_FILL, end_color=PASS_FILL, fill_type="solid")
            status_cell.font = Font(name="Times New Roman", size=10, bold=True, color=PASS_FONT)
        elif student.result_status == "WITHHELD":
            status_cell.fill = PatternFill(start_color=AMBER_FILL, end_color=AMBER_FILL, fill_type="solid")
            status_cell.font = Font(name="Times New Roman", size=10, bold=True, color=AMBER_FONT)
        else:
            status_cell.fill = PatternFill(start_color=FAIL_FILL, end_color=FAIL_FILL, fill_type="solid")
            status_cell.font = Font(name="Times New Roman", size=10, bold=True, color=FAIL_FONT)

        verif_val = "VERIFIED" if student.is_verified else "PENDING"
        verif_cell = ws.cell(row=r, column=verif_col, value=verif_val)
        verif_cell.alignment = align_center
        if student.is_verified:
            verif_cell.font = Font(name="Times New Roman", size=9, bold=True, color=INFO_FONT)
            verif_cell.fill = PatternFill(start_color=INFO_FILL, end_color=INFO_FILL, fill_type="solid")
        else:
            verif_cell.font = Font(name="Times New Roman", size=9, bold=True, color=AMBER_FONT)
            verif_cell.fill = PatternFill(start_color=AMBER_FILL, end_color=AMBER_FILL, fill_type="solid")

        # Apply borders and zebra fill to unstyled cells
        for c in range(1, total_cols + 1):
            cl = ws.cell(row=r, column=c)
            cl.border = table_border
            if not cl.font:
                cl.font = font_data
            if not cl.fill or cl.fill.fill_type is None:
                cl.fill = row_fill

    # Summary Total Row at the bottom
    if matrix.rows:
        summary_row = last_data_row + 1
        ws.row_dimensions[summary_row].height = 24

        # Merge leading label columns A..F
        ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=6)
        lbl_cell = ws.cell(row=summary_row, column=1, value="INSTITUTIONAL TOTALS & SUMMARY")
        lbl_cell.font = font_summary_total
        lbl_cell.alignment = align_center
        lbl_cell.fill = fill_summary_total

        # Subject columns pass count formulas
        for s_idx in range(num_subjects):
            c_idx = sub_col_start + s_idx
            c_letter = get_column_letter(c_idx)
            sub_pass_fml = (
                f'=COUNTIF({c_letter}{data_start_row}:{c_letter}{last_data_row}, "O") + '
                f'COUNTIF({c_letter}{data_start_row}:{c_letter}{last_data_row}, "A+") + '
                f'COUNTIF({c_letter}{data_start_row}:{c_letter}{last_data_row}, "A") + '
                f'COUNTIF({c_letter}{data_start_row}:{c_letter}{last_data_row}, "B+") + '
                f'COUNTIF({c_letter}{data_start_row}:{c_letter}{last_data_row}, "B") + '
                f'COUNTIF({c_letter}{data_start_row}:{c_letter}{last_data_row}, "C")'
            )
            sc = ws.cell(row=summary_row, column=c_idx, value=sub_pass_fml)
            sc.font = font_summary_total
            sc.alignment = align_center
            sc.fill = fill_summary_total

        curr_c_letter = get_column_letter(curr_arr_col)
        sc_curr = ws.cell(row=summary_row, column=curr_arr_col, value=f"=SUM({curr_c_letter}{data_start_row}:{curr_c_letter}{last_data_row})")
        sc_curr.font = font_summary_total
        sc_curr.alignment = align_center
        sc_curr.fill = fill_summary_total

        stand_c_letter = get_column_letter(stand_arr_col)
        sc_stand = ws.cell(row=summary_row, column=stand_arr_col, value=f"=SUM({stand_c_letter}{data_start_row}:{stand_c_letter}{last_data_row})")
        sc_stand.font = font_summary_total
        sc_stand.alignment = align_center
        sc_stand.fill = fill_summary_total

        gpa_c_letter = get_column_letter(gpa_col)
        sc_gpa = ws.cell(row=summary_row, column=gpa_col, value=f"=AVERAGE({gpa_c_letter}{data_start_row}:{gpa_c_letter}{last_data_row})")
        sc_gpa.font = font_summary_total
        sc_gpa.number_format = "0.00"
        sc_gpa.alignment = align_center
        sc_gpa.fill = fill_summary_total

        status_c_letter = get_column_letter(status_col)
        status_summary_fml = (
            f'=COUNTIF({status_c_letter}{data_start_row}:{status_c_letter}{last_data_row}, "PASS") & " Pass / " & '
            f'COUNTIF({status_c_letter}{data_start_row}:{status_c_letter}{last_data_row}, "FAIL") & " Fail"'
        )
        sc_stat = ws.cell(row=summary_row, column=status_col, value=status_summary_fml)
        sc_stat.font = font_summary_total
        sc_stat.alignment = align_center
        sc_stat.fill = fill_summary_total

        verif_c_letter = get_column_letter(verif_col)
        verif_summary_fml = f'=COUNTIF({verif_c_letter}{data_start_row}:{verif_c_letter}{last_data_row}, "VERIFIED") & " Verified"'
        sc_ver = ws.cell(row=summary_row, column=verif_col, value=verif_summary_fml)
        sc_ver.font = font_summary_total
        sc_ver.alignment = align_center
        sc_ver.fill = fill_summary_total

        for c in range(1, total_cols + 1):
            ws.cell(row=summary_row, column=c).border = summary_border

    # Freeze panes: Row 5 header stays visible vertically, Col A (S.No) & Col B (Register No) stay visible horizontally.
    # Col C (Student Name) and dynamic subject columns scroll smoothly.
    apply_sheet_view_and_freeze(ws, freeze_cell="C6")

    # AutoFilter covers all student records
    filter_end_row = last_data_row if matrix.rows else header_row
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{filter_end_row}"

    # Auto column widths
    apply_auto_column_widths(
        ws,
        header_row=header_row,
        max_data_row=last_data_row,
        min_col=1,
        max_col=total_cols,
        compact_cols=compact_cols
    )

    # Conditional formatting rules for live reactivity
    if matrix.rows:
        apply_conditional_formatting_rules(
            ws,
            start_row=data_start_row,
            end_row=last_data_row,
            start_col=sub_col_start,
            end_col=status_col
        )

    # Print settings
    apply_print_settings(ws, landscape=True, print_title_row=header_row)


def build_summary_sheet(ws, matrix):
    ws.title = "Summary"

    sub_headers = ["S.No", "Subject Code", "Subject Name", "Credits", "Appeared", "Passed", "Failed", "Pass %", "O", "A+", "A", "B+", "B", "C", "RA"]
    total_cols = len(sub_headers)
    last_col_letter = get_column_letter(total_cols)

    # Title Banner — dynamically merged across all 15 columns
    ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = f"PERFORMANCE SUMMARY & GRADE DISTRIBUTION — {matrix.session_code}"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_header
    ws["A1"].alignment = align_center
    ws.row_dimensions[1].height = 26

    # Overall Metrics KPI Cards (Rows 3-4)
    kpis = [
        ("Total Candidates", matrix.total_students),
        ("Total Passed", matrix.passed_students),
        ("Total Failed", matrix.failed_students),
        ("Overall Pass %", f"{matrix.pass_percentage}%"),
        ("Average GPA", matrix.average_gpa)
    ]

    for idx, (label, val) in enumerate(kpis, 1):
        cell_lbl = ws.cell(row=3, column=idx, value=label)
        cell_lbl.font = font_kpi_label
        cell_lbl.fill = fill_kpi
        cell_lbl.alignment = align_center
        cell_lbl.border = table_border

        cell_val = ws.cell(row=4, column=idx, value=val)
        cell_val.font = font_kpi_num
        cell_val.fill = fill_kpi
        cell_val.alignment = align_center
        cell_val.border = table_border

    ws.row_dimensions[3].height = 18
    ws.row_dimensions[4].height = 26
    ws.row_dimensions[5].height = 10  # Spacer
    ws.row_dimensions[6].height = 10  # Spacer

    # Subject-wise Breakdown Table (Row 7)
    header_row = 7
    ws.row_dimensions[header_row].height = 28
    for c_idx, h in enumerate(sub_headers, 1):
        cell = ws.cell(row=header_row, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = table_border

    data_start_row = 8
    last_data_row = data_start_row - 1

    # Compute Subject Breakdown Data
    for i, sub in enumerate(matrix.columns, 1):
        r = data_start_row + i - 1
        last_data_row = r
        ws.row_dimensions[r].height = 20
        row_fill = fill_zebra if i % 2 == 0 else fill_white

        appeared = 0
        passed = 0
        failed = 0
        grade_counts = {g: 0 for g in ["O", "A+", "A", "B+", "B", "C", "RA"]}

        for student in matrix.rows:
            gr = student.grades.get(sub.code)
            if gr and gr != "-":
                appeared += 1
                if gr in PASSING_GRADES:
                    passed += 1
                else:
                    failed += 1
                if gr in grade_counts:
                    grade_counts[gr] += 1

        # Pass percentage as formula: =IF(E8>0, F8/E8, 0)
        pass_pct_fml = f"=IF(E{r}>0, F{r}/E{r}, 0)"

        row_vals = [
            (i, align_center, None),
            (sub.code, align_center, None),
            (sub.name, align_left, None),
            (sub.credits, align_center, None),
            (appeared, align_center, None),
            (passed, align_center, None),
            (failed, align_center, None),
            (pass_pct_fml, align_center, "0.0%"),
            (grade_counts["O"], align_center, None),
            (grade_counts["A+"], align_center, None),
            (grade_counts["A"], align_center, None),
            (grade_counts["B+"], align_center, None),
            (grade_counts["B"], align_center, None),
            (grade_counts["C"], align_center, None),
            (grade_counts["RA"], align_center, None),
        ]

        for c_idx, (v, aln, num_fmt) in enumerate(row_vals, 1):
            cl = ws.cell(row=r, column=c_idx, value=v)
            cl.font = font_data
            cl.alignment = aln
            cl.fill = row_fill
            cl.border = table_border
            if num_fmt:
                cl.number_format = num_fmt

    # Summary Total Row at bottom
    if matrix.columns:
        summary_row = last_data_row + 1
        ws.row_dimensions[summary_row].height = 24

        ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=3)
        lbl = ws.cell(row=summary_row, column=1, value="OVERALL TOTAL")
        lbl.font = font_summary_total
        lbl.alignment = align_center
        lbl.fill = fill_summary_total

        # Total Credits: =SUM(D8:D{last})
        tot_cred = ws.cell(row=summary_row, column=4, value=f"=SUM(D{data_start_row}:D{last_data_row})")
        tot_cred.font = font_summary_total
        tot_cred.alignment = align_center
        tot_cred.fill = fill_summary_total

        # Total Appeared: =SUM(E8:E{last})
        tot_app = ws.cell(row=summary_row, column=5, value=f"=SUM(E{data_start_row}:E{last_data_row})")
        tot_app.font = font_summary_total
        tot_app.alignment = align_center
        tot_app.fill = fill_summary_total

        # Total Passed: =SUM(F8:F{last})
        tot_pass = ws.cell(row=summary_row, column=6, value=f"=SUM(F{data_start_row}:F{last_data_row})")
        tot_pass.font = font_summary_total
        tot_pass.alignment = align_center
        tot_pass.fill = fill_summary_total

        # Total Failed: =SUM(G8:G{last})
        tot_fail = ws.cell(row=summary_row, column=7, value=f"=SUM(G{data_start_row}:G{last_data_row})")
        tot_fail.font = font_summary_total
        tot_fail.alignment = align_center
        tot_fail.fill = fill_summary_total

        # Overall Pass %: =IF(E{tot}>0, F{tot}/E{tot}, 0)
        tot_pct = ws.cell(row=summary_row, column=8, value=f"=IF(E{summary_row}>0, F{summary_row}/E{summary_row}, 0)")
        tot_pct.font = font_summary_total
        tot_pct.alignment = align_center
        tot_pct.number_format = "0.0%"
        tot_pct.fill = fill_summary_total

        # Grade counts sums: O (col 9) through RA (col 15)
        for c in range(9, 16):
            c_letter = get_column_letter(c)
            gr_cell = ws.cell(row=summary_row, column=c, value=f"=SUM({c_letter}{data_start_row}:{c_letter}{last_data_row})")
            gr_cell.font = font_summary_total
            gr_cell.alignment = align_center
            gr_cell.fill = fill_summary_total

        for c in range(1, total_cols + 1):
            ws.cell(row=summary_row, column=c).border = summary_border

    # Freeze panes: keep row 7 headers visible vertically, S.No & Subject Code visible horizontally
    apply_sheet_view_and_freeze(ws, freeze_cell="C8")

    # AutoFilter covers subject breakdown table
    filter_end_row = last_data_row if matrix.columns else header_row
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{filter_end_row}"

    # Auto column widths with compact grade counts
    compact_grades = list(range(9, 16))
    apply_auto_column_widths(
        ws,
        header_row=header_row,
        max_data_row=last_data_row,
        min_col=1,
        max_col=total_cols,
        compact_cols=compact_grades
    )

    # Print setup
    apply_print_settings(ws, landscape=True, print_title_row=header_row)


def build_arrear_analysis_sheet(ws, matrix):
    ws.title = "Arrear Analysis"

    headers = ["S.No", "Register No", "Student Name", "Department", "Current Arrears", "Standing Arrears", "Failed Subject Codes"]
    total_cols = len(headers)
    last_col_letter = get_column_letter(total_cols)

    # Title Banner — dynamically merged across all 7 columns
    ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = f"ARREAR ANALYSIS & STANDING ARREARS REPORT — {matrix.session_code}"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_header
    ws["A1"].alignment = align_center
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 10  # Spacer

    header_row = 3
    ws.row_dimensions[header_row].height = 28
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = table_border

    arrear_students = [s for s in matrix.rows if s.arrears_count > 0 or s.standing_arrears_count > 0]
    data_start_row = 4
    last_data_row = data_start_row - 1

    for i, s in enumerate(arrear_students, 1):
        r = data_start_row + i - 1
        last_data_row = r
        ws.row_dimensions[r].height = 22
        row_fill = fill_zebra if i % 2 == 0 else fill_white

        failed_subs = [code for code, status in s.subject_statuses.items() if status == "RA"]

        reg_str = str(s.register_number).strip()
        reg_val = int(reg_str) if reg_str.isdigit() and not reg_str.startswith("0") else reg_str
        vals = [
            (i, align_center, None),
            (reg_val, align_center, "0" if isinstance(reg_val, int) else None),
            (s.student_name, align_left, None),
            (s.department_code, align_center, None),
            (s.arrears_count, align_center, None),
            (s.standing_arrears_count, align_center, None),
            (", ".join(failed_subs) if failed_subs else "-", align_left_wrap, None)
        ]

        for c_idx, (v, aln, num_fmt) in enumerate(vals, 1):
            cl = ws.cell(row=r, column=c_idx, value=v)
            cl.font = font_data
            cl.alignment = aln
            cl.fill = row_fill
            cl.border = table_border
            if num_fmt:
                cl.number_format = num_fmt

            # Highlight arrear numbers
            if c_idx in [5, 6] and isinstance(v, int) and v > 0:
                cl.font = Font(name="Times New Roman", size=10, bold=True, color=FAIL_FONT)
                cl.fill = PatternFill(start_color=FAIL_FILL, end_color=FAIL_FILL, fill_type="solid")

    # Total Summary Row at bottom
    if arrear_students:
        summary_row = last_data_row + 1
        ws.row_dimensions[summary_row].height = 24

        ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=4)
        lbl = ws.cell(row=summary_row, column=1, value=f"TOTAL CANDIDATES WITH ARREARS: {len(arrear_students)}")
        lbl.font = font_summary_total
        lbl.alignment = align_center
        lbl.fill = fill_summary_total

        tot_curr = ws.cell(row=summary_row, column=5, value=f"=SUM(E{data_start_row}:E{last_data_row})")
        tot_curr.font = font_summary_total
        tot_curr.alignment = align_center
        tot_curr.fill = fill_summary_total

        tot_stand = ws.cell(row=summary_row, column=6, value=f"=SUM(F{data_start_row}:F{last_data_row})")
        tot_stand.font = font_summary_total
        tot_stand.alignment = align_center
        tot_stand.fill = fill_summary_total

        cl_g = ws.cell(row=summary_row, column=7, value="-")
        cl_g.font = font_summary_total
        cl_g.alignment = align_center
        cl_g.fill = fill_summary_total

        for c in range(1, total_cols + 1):
            ws.cell(row=summary_row, column=c).border = summary_border

    # Freeze panes: row 3 headers visible vertically, Col A & B visible horizontally
    apply_sheet_view_and_freeze(ws, freeze_cell="C4")

    # AutoFilter
    filter_end_row = last_data_row if arrear_students else header_row
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{filter_end_row}"

    # Auto column widths
    apply_auto_column_widths(ws, header_row=header_row, max_data_row=last_data_row, min_col=1, max_col=total_cols)

    # Print settings
    apply_print_settings(ws, landscape=True, print_title_row=header_row)


def build_processing_log_sheet(ws, db, matrix):
    ws.title = "Processing Log"

    headers = ["Log ID", "Register No", "Student Name", "Stage", "Status", "Processing Details", "Timestamp"]
    total_cols = len(headers)
    last_col_letter = get_column_letter(total_cols)

    # Title Banner — dynamically merged across all 7 columns
    ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = f"RESULT RETRIEVAL & EXTRACTION AUDIT LOG — {matrix.session_code}"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_header
    ws["A1"].alignment = align_center
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 10  # Spacer

    header_row = 3
    ws.row_dimensions[header_row].height = 28
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = table_border

    logs = db.query(AutomationJobLog).order_by(AutomationJobLog.id.desc()).limit(200).all()
    data_start_row = 4
    last_data_row = data_start_row - 1

    for i, l in enumerate(logs, 1):
        r = data_start_row + i - 1
        last_data_row = r
        ws.row_dimensions[r].height = 20
        row_fill = fill_zebra if i % 2 == 0 else fill_white

        reg_str = str(l.register_number).strip() if l.register_number else "-"
        reg_val = int(reg_str) if reg_str.isdigit() and not reg_str.startswith("0") else reg_str
        vals = [
            (l.id, align_center, None),
            (reg_val, align_center, "0" if isinstance(reg_val, int) else None),
            (l.student_name or "-", align_left, None),
            (l.stage, align_center, None),
            (l.status, align_center, None),
            (l.message, align_left_wrap, None),
            (l.timestamp.strftime("%Y-%m-%d %H:%M:%S"), align_center, None)
        ]

        for c_idx, (v, aln, num_fmt) in enumerate(vals, 1):
            cl = ws.cell(row=r, column=c_idx, value=v)
            cl.font = font_data
            cl.alignment = aln
            cl.fill = row_fill
            cl.border = table_border
            if num_fmt:
                cl.number_format = num_fmt

            # Highlight Status Column (Col 5)
            if c_idx == 5:
                st_upper = str(v).upper()
                if "SUCCESS" in st_upper or "COMPLETE" in st_upper:
                    cl.font = Font(name="Times New Roman", size=10, bold=True, color=PASS_FONT)
                    cl.fill = PatternFill(start_color=PASS_FILL, end_color=PASS_FILL, fill_type="solid")
                elif "FAIL" in st_upper or "ERROR" in st_upper:
                    cl.font = Font(name="Times New Roman", size=10, bold=True, color=FAIL_FONT)
                    cl.fill = PatternFill(start_color=FAIL_FILL, end_color=FAIL_FILL, fill_type="solid")
                elif "PEND" in st_upper or "PROCESS" in st_upper or "RUN" in st_upper:
                    cl.font = Font(name="Times New Roman", size=10, bold=True, color=AMBER_FONT)
                    cl.fill = PatternFill(start_color=AMBER_FILL, end_color=AMBER_FILL, fill_type="solid")

    # Summary Total Row
    if logs:
        summary_row = last_data_row + 1
        ws.row_dimensions[summary_row].height = 24

        ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=7)
        lbl = ws.cell(row=summary_row, column=1, value=f"TOTAL LOG ENTRIES RECORDED: {len(logs)}")
        lbl.font = font_summary_total
        lbl.alignment = align_center
        lbl.fill = fill_summary_total

        for c in range(1, total_cols + 1):
            ws.cell(row=summary_row, column=c).border = summary_border

    # Freeze panes: Row 3 headers visible vertically, Col A & B visible horizontally
    apply_sheet_view_and_freeze(ws, freeze_cell="C4")

    # AutoFilter
    filter_end_row = last_data_row if logs else header_row
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{filter_end_row}"

    # Auto column widths
    apply_auto_column_widths(ws, header_row=header_row, max_data_row=last_data_row, min_col=1, max_col=total_cols, max_width_cap=55)

    # Print settings
    apply_print_settings(ws, landscape=True, print_title_row=header_row)


def build_subject_mapping_sheet(ws, matrix):
    ws.title = "Subject Mapping"

    headers = ["S.No", "Subject Code", "Course Title", "Credits", "Department"]
    total_cols = len(headers)
    last_col_letter = get_column_letter(total_cols)

    # Title Banner — dynamically merged across all 5 columns
    ws.merge_cells(f"A1:{last_col_letter}1")
    ws["A1"] = f"COURSE SYLLABUS & SUBJECT DEFINITIONS — {matrix.session_code}"
    ws["A1"].font = font_title
    ws["A1"].fill = fill_header
    ws["A1"].alignment = align_center
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 10  # Spacer

    header_row = 3
    ws.row_dimensions[header_row].height = 28
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = table_border

    data_start_row = 4
    last_data_row = data_start_row - 1

    for i, sub in enumerate(matrix.columns, 1):
        r = data_start_row + i - 1
        last_data_row = r
        ws.row_dimensions[r].height = 20
        row_fill = fill_zebra if i % 2 == 0 else fill_white

        vals = [
            (i, align_center),
            (sub.code, align_center),
            (sub.name, align_left),
            (sub.credits, align_center),
            (matrix.department_code or "ALL", align_center)
        ]
        for c_idx, (v, aln) in enumerate(vals, 1):
            cl = ws.cell(row=r, column=c_idx, value=v)
            cl.font = font_data
            cl.alignment = aln
            cl.fill = row_fill
            cl.border = table_border

    # Summary Total Row
    if matrix.columns:
        summary_row = last_data_row + 1
        ws.row_dimensions[summary_row].height = 24

        ws.merge_cells(start_row=summary_row, start_column=1, end_row=summary_row, end_column=3)
        lbl = ws.cell(row=summary_row, column=1, value=f"TOTAL SUBJECTS: {len(matrix.columns)}")
        lbl.font = font_summary_total
        lbl.alignment = align_center
        lbl.fill = fill_summary_total

        tot_cred = ws.cell(row=summary_row, column=4, value=f"=SUM(D{data_start_row}:D{last_data_row})")
        tot_cred.font = font_summary_total
        tot_cred.alignment = align_center
        tot_cred.fill = fill_summary_total

        cl_dept = ws.cell(row=summary_row, column=5, value=matrix.department_code or "ALL")
        cl_dept.font = font_summary_total
        cl_dept.alignment = align_center
        cl_dept.fill = fill_summary_total

        for c in range(1, total_cols + 1):
            ws.cell(row=summary_row, column=c).border = summary_border

    # Freeze panes: Row 3 headers visible vertically, Col A & B visible horizontally
    apply_sheet_view_and_freeze(ws, freeze_cell="C4")

    # AutoFilter
    filter_end_row = last_data_row if matrix.columns else header_row
    ws.auto_filter.ref = f"A{header_row}:{last_col_letter}{filter_end_row}"

    # Auto column widths
    apply_auto_column_widths(ws, header_row=header_row, max_data_row=last_data_row, min_col=1, max_col=total_cols)

    # Print settings
    apply_print_settings(ws, landscape=False, print_title_row=header_row)


def validate_generated_workbook(file_path: str, expected_sheets: List[str]) -> Tuple[bool, str]:
    """
    Reopens and verifies workbook structural integrity, sheet presence, cell content,
    freeze panes, and accessibility.
    """
    try:
        wb = openpyxl.load_workbook(file_path, data_only=False)
        sheet_names = wb.sheetnames

        for s in expected_sheets:
            if s not in sheet_names:
                return False, f"Missing required worksheet: {s}"

        # 1. Verify Result Analysis has data & headers
        ra_sheet = wb["Result Analysis"]
        if ra_sheet.max_row < 5:
            return False, "Result Analysis sheet is empty or missing headers."

        # Verify institutional header and register no column
        if not ra_sheet["A1"].value:
            return False, "Result Analysis title header is empty."
        if ra_sheet["B5"].value != "Register No":
            return False, f"Expected Register No at B5, got {ra_sheet['B5'].value}"

        # 2. Verify Freeze Panes on Result Analysis
        if ra_sheet.freeze_panes != "C6":
            return False, f"Result Analysis freeze panes should be 'C6', got '{ra_sheet.freeze_panes}'"

        # 3. Verify no unintended sheet protection
        for s_name in sheet_names:
            ws = wb[s_name]
            if ws.protection and ws.protection.sheet:
                return False, f"Worksheet '{s_name}' is locked or protected."

        # 4. Verify AutoFilter is configured on all sheets
        for s_name in expected_sheets:
            ws = wb[s_name]
            if not ws.auto_filter or not ws.auto_filter.ref:
                return False, f"Worksheet '{s_name}' missing AutoFilter range."

        return True, "Workbook integrity verified successfully."
    except Exception as e:
        return False, f"Workbook verification failed: {str(e)}"


def inject_ignored_errors_to_workbook_file(file_path: str) -> None:
    """
    Injects SpreadsheetML <ignoredErrors> elements into the saved .xlsx ZIP archive
    to permanently remove Excel's 'Number Stored as Text' green triangle warnings on register numbers
    while preserving the strict text format '@'.
    """
    import zipfile
    import os

    temp_path = file_path + ".tmp"
    with zipfile.ZipFile(file_path, "r") as zin, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet") and item.filename.endswith(".xml"):
                xml_str = data.decode("utf-8")
                if "</worksheet>" in xml_str and "<ignoredErrors>" not in xml_str:
                    ignored_block = '<ignoredErrors><ignoredError sqref="A4:C5000" numberStoredAsText="1"/></ignoredErrors></worksheet>'
                    xml_str = xml_str.replace("</worksheet>", ignored_block)
                    data = xml_str.encode("utf-8")
            zout.writestr(item, data)

    os.replace(temp_path, file_path)


def generate_enterprise_excel_report(
    db: Session,
    session_id: int,
    department_code: Optional[str] = None,
    batch: Optional[str] = None,
    semester: Optional[int] = None,
    generated_by: str = "SYSTEM"
) -> ExcelReport:
    """
    Generates, formats, verifies, and records an enterprise examination analysis report.
    """
    matrix = get_result_matrix(db, session_id, department_code, batch, semester)
    session_obj = db.query(ExamSession).filter(ExamSession.id == session_id).first()

    wb = openpyxl.Workbook()
    # Remove default sheet
    default_sheet = wb.active

    # 1. Result Analysis Sheet
    ws_ra = wb.create_sheet("Result Analysis")
    build_result_analysis_sheet(ws_ra, matrix)

    # 2. Summary Sheet
    ws_sum = wb.create_sheet("Summary")
    build_summary_sheet(ws_sum, matrix)

    # 3. Arrear Analysis Sheet
    ws_arr = wb.create_sheet("Arrear Analysis")
    build_arrear_analysis_sheet(ws_arr, matrix)

    # 4. Processing Log Sheet
    ws_log = wb.create_sheet("Processing Log")
    build_processing_log_sheet(ws_log, db, matrix)

    # 5. Subject Mapping Sheet
    ws_sub = wb.create_sheet("Subject Mapping")
    build_subject_mapping_sheet(ws_sub, matrix)

    # Remove initial empty sheet
    wb.remove(default_sheet)

    # Strictly enforce Times New Roman on every cell across all worksheets
    for ws_each in wb.worksheets:
        for row in ws_each.iter_rows():
            for cell in row:
                if cell.font:
                    if cell.font.name != "Times New Roman":
                        cell.font = Font(
                            name="Times New Roman",
                            size=cell.font.size,
                            bold=cell.font.bold,
                            italic=cell.font.italic,
                            vertAlign=cell.font.vertAlign,
                            underline=cell.font.underline,
                            strike=cell.font.strike,
                            color=cell.font.color
                        )
                else:
                    cell.font = font_data

    # File naming & storage
    timestamp_str = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    dept_str = department_code or "ALL"
    filename = f"EXAM_RESULT_{matrix.session_code}_{dept_str}_{timestamp_str}.xlsx"
    file_path = str(settings.REPORTS_DIR / filename)

    wb.save(file_path)
    inject_ignored_errors_to_workbook_file(file_path)

    # Validation
    expected_sheets = ["Result Analysis", "Summary", "Arrear Analysis", "Processing Log", "Subject Mapping"]
    is_valid, validation_notes = validate_generated_workbook(file_path, expected_sheets)

    # Compute SHA256 & file size
    sha256_hash = hashlib.sha256()
    with open(file_path, "rb") as f:
        for block in iter(lambda: f.read(4096), b""):
            sha256_hash.update(block)
    file_hash = sha256_hash.hexdigest()
    file_size = os.path.getsize(file_path)

    # Store in database
    report = ExcelReport(
        session_id=session_id,
        department_code=department_code,
        batch=batch,
        semester=semester,
        report_name=filename,
        file_path=file_path,
        file_size=file_size,
        file_hash=file_hash,
        total_students=matrix.total_students,
        total_subjects=len(matrix.columns),
        pass_percentage=matrix.pass_percentage,
        is_valid=is_valid,
        validation_notes=validation_notes,
        generated_by=generated_by
    )
    db.add(report)
    db.commit()
    db.refresh(report)

    return report
