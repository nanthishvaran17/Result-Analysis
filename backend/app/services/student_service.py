import io
from typing import Optional, List, Tuple
import pandas as pd
from sqlalchemy.orm import Session
from sqlalchemy import or_
from app.models.student import Student, Department
from app.schemas.student import StudentCreate, StudentUpdate, StudentImportSummary, StudentImportErrorItem

def get_students(
    db: Session,
    department_code: Optional[str] = None,
    batch: Optional[str] = None,
    semester: Optional[int] = None,
    search: Optional[str] = None,
    skip: int = 0,
    limit: int = 50
) -> Tuple[List[Student], int]:
    query = db.query(Student)
    if department_code:
        query = query.filter(Student.department_code == department_code)
    if batch:
        query = query.filter(Student.batch == batch)
    if semester:
        query = query.filter(Student.current_semester == semester)
    if search:
        search_fmt = f"%{search.strip()}%"
        query = query.filter(
            or_(
                Student.register_number.ilike(search_fmt),
                Student.name.ilike(search_fmt),
                Student.email.ilike(search_fmt)
            )
        )
    total = query.count()
    students = query.order_by(Student.register_number.asc()).offset(skip).limit(limit).all()
    return students, total

def get_student_by_id(db: Session, student_id: int) -> Optional[Student]:
    return db.query(Student).filter(Student.id == student_id).first()

def get_student_by_reg_no(db: Session, reg_no: str) -> Optional[Student]:
    return db.query(Student).filter(Student.register_number == reg_no.strip()).first()

def create_student(db: Session, student_in: StudentCreate) -> Student:
    reg_no = student_in.register_number.strip().upper()
    existing = get_student_by_reg_no(db, reg_no)
    if existing:
        raise ValueError(f"Student with Register Number {reg_no} already exists.")
    
    # Ensure department exists
    dept = db.query(Department).filter(Department.code == student_in.department_code.upper()).first()
    if not dept:
        dept = Department(code=student_in.department_code.upper(), name=f"Dept of {student_in.department_code.upper()}")
        db.add(dept)
        db.flush()

    student = Student(
        register_number=reg_no,
        name=student_in.name.strip(),
        department_code=student_in.department_code.upper(),
        batch=student_in.batch.strip(),
        current_semester=student_in.current_semester,
        section=student_in.section.strip().upper() if student_in.section else "A",
        email=student_in.email,
        is_active=True
    )
    db.add(student)
    db.commit()
    db.refresh(student)
    return student

def update_student(db: Session, student_id: int, student_in: StudentUpdate) -> Optional[Student]:
    student = get_student_by_id(db, student_id)
    if not student:
        return None
    if student_in.name is not None:
        student.name = student_in.name.strip()
    if student_in.department_code is not None:
        student.department_code = student_in.department_code.upper()
    if student_in.batch is not None:
        student.batch = student_in.batch.strip()
    if student_in.current_semester is not None:
        student.current_semester = student_in.current_semester
    if student_in.section is not None:
        student.section = student_in.section.strip().upper()
    if student_in.email is not None:
        student.email = student_in.email
    if student_in.is_active is not None:
        student.is_active = student_in.is_active
    db.commit()
    db.refresh(student)
    return student

def import_students_from_file(db: Session, file_contents: bytes, filename: str) -> StudentImportSummary:
    summary = StudentImportSummary(
        total_rows=0,
        imported_count=0,
        duplicate_count=0,
        error_count=0,
        errors=[]
    )
    
    try:
        if filename.endswith(".csv"):
            df = pd.read_csv(io.BytesIO(file_contents), dtype=str)
        elif filename.endswith((".xlsx", ".xls")):
            df = pd.read_excel(io.BytesIO(file_contents), dtype=str)
        else:
            raise ValueError("Unsupported file format. Please upload CSV or Excel (.xlsx, .xls) file.")
    except Exception as e:
        summary.error_count += 1
        summary.errors.append(StudentImportErrorItem(row=0, error=f"File parsing error: {str(e)}"))
        return summary

    # Normalize column names: strip whitespace and lowercase
    col_map = {col: col.strip().lower().replace(" ", "_") for col in df.columns}
    df.rename(columns=col_map, inplace=True)

    required_cols = ["register_number", "name", "department"]
    missing = [c for c in required_cols if not any(c in col for col in df.columns)]
    
    # Try alternate column names if standard not found
    reg_col = next((c for c in df.columns if ("reg" in c or "roll" in c) and "controller" not in c and "instruction" not in c), None)
    name_col = next((c for c in df.columns if "name" in c and "instruction" not in c), None)
    dept_col = next((c for c in df.columns if ("dept" in c or "branch" in c or "department" in c) and "instruction" not in c), None)
    batch_col = next((c for c in df.columns if ("batch" in c or "year" in c) and "instruction" not in c), None)
    sem_col = next((c for c in df.columns if "sem" in c and "instruction" not in c), None)
    sec_col = next((c for c in df.columns if "sec" in c and "instruction" not in c), None)
    email_col = next((c for c in df.columns if ("email" in c or "mail" in c) and "instruction" not in c), None)

    if not reg_col or not name_col or not dept_col:
        summary.error_count += 1
        summary.errors.append(StudentImportErrorItem(
            row=0,
            error=f"Missing required columns. Required at least: Register Number, Name, Department. Found: {list(df.columns)}"
        ))
        return summary

    summary.total_rows = len(df)
    existing_reg_nos = {r[0] for r in db.query(Student.register_number).all()}
    departments_cache = {d.code: d for d in db.query(Department).all()}

    new_students = []
    
    for idx, row in df.iterrows():
        row_num = idx + 2  # 1-indexed header + row
        reg_val = str(row[reg_col]).strip().upper() if pd.notna(row[reg_col]) else ""
        name_val = str(row[name_col]).strip() if pd.notna(row[name_col]) else ""
        dept_val = str(row[dept_col]).strip().upper() if pd.notna(row[dept_col]) else "CSE"
        batch_val = str(row[batch_col]).strip() if batch_col and pd.notna(row[batch_col]) else "2022-2026"
        sem_val = int(row[sem_col]) if sem_col and pd.notna(row[sem_col]) and str(row[sem_col]).strip().isdigit() else 5
        sec_val = str(row[sec_col]).strip().upper() if sec_col and pd.notna(row[sec_col]) else "A"
        email_val = str(row[email_col]).strip() if email_col and pd.notna(row[email_col]) else None

        # Safely ignore trailing empty rows from templates
        if (not reg_val or reg_val == "NAN") and (not name_val or name_val == "NAN"):
            continue

        if not reg_val or reg_val == "NAN":
            summary.error_count += 1
            summary.errors.append(StudentImportErrorItem(row=row_num, register_number=reg_val, error="Register Number is empty"))
            continue

        if reg_val in existing_reg_nos:
            summary.duplicate_count += 1
            summary.errors.append(StudentImportErrorItem(row=row_num, register_number=reg_val, error="Duplicate Register Number already registered"))
            continue

        if not name_val or name_val == "NAN":
            summary.error_count += 1
            summary.errors.append(StudentImportErrorItem(row=row_num, register_number=reg_val, error="Student name is missing"))
            continue

        # Ensure department exists in cache / DB
        if dept_val not in departments_cache:
            dept_obj = Department(code=dept_val, name=f"Department of {dept_val}")
            db.add(dept_obj)
            db.flush()
            departments_cache[dept_val] = dept_obj

        student = Student(
            register_number=reg_val,
            name=name_val,
            department_code=dept_val,
            batch=batch_val,
            current_semester=sem_val,
            section=sec_val,
            email=email_val,
            is_active=True
        )
        new_students.append(student)
        existing_reg_nos.add(reg_val)
        summary.imported_count += 1

    if new_students:
        db.add_all(new_students)
        db.commit()

    return summary

def bulk_create_students(db: Session, students_in: List[StudentCreate]) -> Dict[str, Any]:
    """
    Directly persists a batch of student records entered via the live UI spreadsheet grid.
    Automatically handles existing duplicates and ensures academic departments exist.
    """
    departments_cache = {d.code: d for d in db.query(Department).all()}
    existing_reg_nos = {r[0] for r in db.query(Student.register_number).all()}
    
    created_list: List[Student] = []
    skipped_list: List[Dict[str, str]] = []

    for item in students_in:
        reg_no = (item.register_number or "").strip().upper()
        name = (item.name or "").strip()
        dept = (item.department_code or "CSE").strip().upper()
        batch = (item.batch or "2022-2026").strip()
        semester = item.current_semester or 5
        section = (item.section or "A").strip().upper()
        email = (item.email or f"{reg_no.lower()}@college.edu").strip()

        if not reg_no or not name:
            skipped_list.append({"register_number": reg_no or "N/A", "reason": "Missing Register Number or Name"})
            continue

        if reg_no in existing_reg_nos:
            skipped_list.append({"register_number": reg_no, "reason": "Already exists in database"})
            continue

        if dept not in departments_cache:
            dept_obj = Department(code=dept, name=f"Department of {dept}")
            db.add(dept_obj)
            db.flush()
            departments_cache[dept] = dept_obj

        s = Student(
            register_number=reg_no,
            name=name,
            department_code=dept,
            batch=batch,
            current_semester=semester,
            section=section,
            email=email,
            is_active=True
        )
        created_list.append(s)
        existing_reg_nos.add(reg_no)

    if created_list:
        db.add_all(created_list)
        db.commit()
        for s in created_list:
            db.refresh(s)

    return {
        "created_count": len(created_list),
        "skipped_count": len(skipped_list),
        "skipped": skipped_list,
        "created_ids": [s.id for s in created_list]
    }

def generate_roster_template() -> io.BytesIO:
    """
    Generates an official, academic-grade Excel (.xlsx) template for student cohort enrollment.
    Strictly uses Times New Roman typography, crisp all-border formatting on every cell, and a clean,
    sober institutional design.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Roster Template"
    if ws.views.sheetView:
        ws.views.sheetView[0].showGridLines = True
    else:
        ws.sheet_view.showGridLines = True

    # Styling tokens - strictly Times New Roman with clean institutional colors
    font_title = Font(name="Times New Roman", size=13, bold=True, color="000000")
    font_sub = Font(name="Times New Roman", size=10, italic=True, color="334155")
    font_hdr = Font(name="Times New Roman", size=11, bold=True, color="000000")
    font_row = Font(name="Times New Roman", size=11, color="000000")

    fill_title = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_sub = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_hdr = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Side(border_style="thin", color="000000")
    all_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    def style_merged_range(sheet, min_c, min_r, max_c, max_r, font, fill, border, alignment):
        for r in range(min_r, max_r + 1):
            for c in range(min_c, max_c + 1):
                cell = sheet.cell(row=r, column=c)
                if font: cell.font = font
                if fill: cell.fill = fill
                if border: cell.border = border
                if alignment: cell.alignment = alignment

    # Row 1: Institutional Header
    ws.merge_cells("A1:G1")
    ws["A1"] = "OFFICE OF THE CONTROLLER OF EXAMINATIONS — STUDENT ROSTER TEMPLATE"
    style_merged_range(ws, 1, 1, 7, 1, font_title, fill_title, all_border, Alignment(horizontal="center", vertical="center"))

    # Row 2: Official Instructions
    ws.merge_cells("A2:G2")
    ws["A2"] = "Instructions: Fill student records in rows 4 onwards. 'Register Number', 'Name', 'Department Code', 'Batch', and 'Semester' are mandatory."
    style_merged_range(ws, 1, 2, 7, 2, font_sub, fill_sub, all_border, Alignment(horizontal="center", vertical="center"))

    # Row 3: Column Headers
    headers = [
        "Register Number", "Student Name", "Department Code", "Batch", "Current Semester", "Section", "Email Address"
    ]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.border = all_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Sample Guidance Rows
    samples = [
        ("710022104051", "Aravindhan M", "CSE", "2022-2026", 5, "A", "710022104051@college.edu"),
        ("710022104052", "Bhavani S", "CSE", "2022-2026", 5, "B", "710022104052@college.edu"),
        ("710022106001", "Chandran P", "ECE", "2023-2027", 3, "A", "710022106001@college.edu"),
    ]

    for r_idx, row_vals in enumerate(samples, 4):
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_row
            cell.fill = fill_white
            cell.border = all_border
            cell.alignment = Alignment(horizontal="center" if c_idx in [1, 3, 4, 5, 6] else "left", vertical="center")

    # Pre-format entry rows 7 to 35 with all borders and Times New Roman
    for r_idx in range(7, 36):
        for c_idx in range(1, 8):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = font_row
            cell.fill = fill_white
            cell.border = all_border
            cell.alignment = Alignment(horizontal="center" if c_idx in [1, 3, 4, 5, 6] else "left", vertical="center")

    # Column Widths
    col_widths = [20, 28, 18, 16, 18, 12, 32]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Row Heights
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 28
    for r in range(4, 36):
        ws.row_dimensions[r].height = 22

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_students_excel(students: List[Student]) -> io.BytesIO:
    """
    Exports registered students into a high-standard, academic Excel sheet.
    Strictly uses Times New Roman font, clean professional headers, and all-border cell formatting.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Directory"
    if ws.views.sheetView:
        ws.views.sheetView[0].showGridLines = True
    else:
        ws.sheet_view.showGridLines = True

    font_title = Font(name="Times New Roman", size=13, bold=True, color="000000")
    font_hdr = Font(name="Times New Roman", size=11, bold=True, color="000000")
    font_row = Font(name="Times New Roman", size=11, color="000000")

    fill_title = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_hdr = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    thin_border = Side(border_style="thin", color="000000")
    all_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    def style_merged_range(sheet, min_c, min_r, max_c, max_r, font, fill, border, alignment):
        for r in range(min_r, max_r + 1):
            for c in range(min_c, max_c + 1):
                cell = sheet.cell(row=r, column=c)
                if font: cell.font = font
                if fill: cell.fill = fill
                if border: cell.border = border
                if alignment: cell.alignment = alignment

    ws.merge_cells("A1:G1")
    ws["A1"] = f"OFFICE OF THE CONTROLLER OF EXAMINATIONS — REGISTERED CANDIDATE ROSTER ({len(students)} Records)"
    style_merged_range(ws, 1, 1, 7, 1, font_title, fill_title, all_border, Alignment(horizontal="center", vertical="center"))

    headers = ["Register Number", "Student Name", "Department", "Batch", "Semester", "Section", "Email"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=2, column=col_idx, value=h)
        cell.font = font_hdr
        cell.fill = fill_hdr
        cell.border = all_border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, s in enumerate(students, 3):
        vals = [s.register_number, s.name, s.department_code, s.batch, s.current_semester, s.section, s.email or ""]
        for c_idx, val in enumerate(vals, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_row
            cell.fill = fill_white
            cell.border = all_border
            cell.alignment = Alignment(horizontal="center" if c_idx in [1, 3, 4, 5, 6] else "left", vertical="center")

    col_widths = [20, 28, 16, 16, 14, 12, 32]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 26
    for r in range(3, len(students) + 3):
        ws.row_dimensions[r].height = 22

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
