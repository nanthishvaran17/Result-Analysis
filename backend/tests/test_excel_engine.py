import os
import openpyxl
from app.core.database import SessionLocal
from app.services.excel_service import generate_enterprise_excel_report, validate_generated_workbook

def test_enterprise_excel_engine_validation():
    db = SessionLocal()
    try:
        report = generate_enterprise_excel_report(
            db=db,
            session_id=1,
            department_code="CSE",
            generated_by="TEST_RUNNER"
        )
        assert report is not None
        assert report.is_valid is True
        assert os.path.exists(report.file_path)

        # Validate sheet names and structure
        wb = openpyxl.load_workbook(report.file_path, data_only=True)
        sheet_names = wb.sheetnames
        expected_sheets = ["Result Analysis", "Summary", "Arrear Analysis", "Processing Log", "Subject Mapping"]
        for s in expected_sheets:
            assert s in sheet_names

        # Check Result Analysis sheet contents
        ra_sheet = wb["Result Analysis"]
        assert ra_sheet["A1"].value is not None  # Institutional Header
        assert ra_sheet["B5"].value == "Register No"
        assert ra_sheet.max_row >= 10

        # Check self-validation
        valid, msg = validate_generated_workbook(report.file_path, expected_sheets)
        assert valid is True
    finally:
        db.close()
